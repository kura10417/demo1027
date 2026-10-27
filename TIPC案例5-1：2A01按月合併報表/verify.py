import importlib.util, pathlib, io, zipfile
import openpyxl
spec=importlib.util.spec_from_file_location('app',pathlib.Path(__file__).with_name('server.py'))
app=importlib.util.module_from_spec(spec);spec.loader.exec_module(app)
source=pathlib.Path(__file__).parent.parent/'附件'/'CH5-1附件'
def load(p): return app.read_file(p.name,p.read_bytes())
t=load(next(source.glob('輸出*.xlsx')))
files=[load(p) for p in sorted(source.glob('輸入_*.xlsx'))]
choices=[{'id':f['id'],'sheet':s['name'],'year':s['year'],'month':s['month']} for f in files for s in f['sheets']]
req={'template':t['id'],'sheet':t['sheets'][0]['name'],'inputs':choices}
r=app.merge(req);raw=app.RESULTS[r['result']]; out=openpyxl.load_workbook(io.BytesIO(raw),data_only=True).active
assert r['months']==list(range(1,13)),r['months']
assert len(r['warnings'])==2 and all('6641110' in w for w in r['warnings'])
for a in r['audit']:
    total=sum(out.cell(row,a['month']+4).value for row in range(4,81))
    omitted={3:2940,9:16734}.get(a['month'],0)
    assert abs(total+omitted-a['total'])<1e-6,(a,total)
original=zipfile.ZipFile(app.FILES[t['id']]['path']); final=zipfile.ZipFile(io.BytesIO(raw))
assert original.namelist()==final.namelist()
changed=[n for n in original.namelist() if original.read(n)!=final.read(n)]
assert changed==['xl/worksheets/sheet1.xml'],changed
original.close(); final.close()
multi=load(next(source.glob('輸入多月*.xlsx')))
mc=[{'id':multi['id'],'sheet':s['name'],'year':s['year'],'month':s['month']} for s in multi['sheets']]
rr=app.merge(dict(req,inputs=mc)); assert rr['months']==[1,2,3]
try: app.merge(dict(req,inputs=choices+mc))
except ValueError as e: assert '超過一張' in str(e)
else: raise AssertionError('duplicate month not rejected')
assert out['E9'].value==213 and out['E80'].value==4865803
# Rows and month columns can move without coordinate assumptions.
tw=openpyxl.load_workbook(app.FILES[t['id']]['path']); ts=tw.active
for row in range(1,81): ts.cell(row,5).value,ts.cell(row,6).value=ts.cell(row,6).value,ts.cell(row,5).value
for col in range(1,17): ts.cell(9,col).value,ts.cell(10,col).value=ts.cell(10,col).value,ts.cell(9,col).value
b=io.BytesIO();tw.save(b); moved=app.read_file('moved.xlsx',b.getvalue())
mr=app.merge(dict(req,template=moved['id'])); mo=openpyxl.load_workbook(io.BytesIO(app.RESULTS[mr['result']]),data_only=True).active
assert mo['F10'].value==213
print('PASS: 12 months reconciled; 3-sheet input; duplicate protection; moved rows/months; all non-data ZIP parts identical.')
