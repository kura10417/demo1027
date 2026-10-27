import base64, hashlib, io, json, math, os, re, secrets, subprocess, tempfile, threading, webbrowser, zipfile
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from xml.sax.saxutils import escape
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
TEMP = tempfile.TemporaryDirectory(prefix='.tipc-temp-', dir=ROOT)
FILES, RESULTS = {}, {}
TOKEN = secrets.token_urlsafe(24)

def norm(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def period(s):
    text = ' '.join(norm(s.cell(r,c).value) for r in range(1,min(s.max_row,4)+1) for c in range(1,min(s.max_column,16)+1))
    m = re.search(r'(\d{3,4})年\s*(\d{1,2})月', text)
    if not m: m = re.search(r'\((\d{3,4})/(\d{1,2})/', text)
    return (int(m[1]),int(m[2])) if m else (None,None)

def header(s):
    for r in range(1,min(s.max_row,30)+1):
        d = {norm(c.value): c.column for c in s[r] if c.value is not None}
        if '貨物群組代碼' in d and '貨物代碼' in d: return r,d
    raise ValueError('找不到「貨物群組代碼」與「貨物代碼」表頭：'+s.title)

def read_file(name,data):
    ext=Path(name).suffix.lower()
    if ext not in ('.xls','.xlsx'): raise ValueError('僅接受 .xls、.xlsx')
    if len(data)>40*1024*1024: raise ValueError('單檔上限 40 MB')
    fid=secrets.token_hex(12); path=Path(TEMP.name)/(fid+ext); path.write_bytes(data)
    converted=False
    if ext=='.xls':
        lo=Path(os.environ.get('PROGRAMFILES','C:/Program Files'))/'LibreOffice/program/soffice.exe'
        if not lo.exists(): raise ValueError('讀取 xls 需要安裝 LibreOffice，或先以 Excel 另存成 xlsx。')
        profile=(Path(TEMP.name)/('profile-'+fid)).as_uri()
        p=subprocess.run([str(lo),'-env:UserInstallation='+profile,'--headless','--convert-to','xlsx','--outdir',TEMP.name,str(path)],capture_output=True,timeout=90)
        path=path.with_suffix('.xlsx')
        if p.returncode or not path.exists(): raise ValueError('xls 轉換失敗，請另存成 xlsx 後重試。')
        converted=True
    book=openpyxl.load_workbook(io.BytesIO(path.read_bytes()),data_only=False)
    if sum(s.max_row*s.max_column for s in book)>2000000: raise ValueError('工作表範圍過大，請先移除多餘空白列欄。')
    FILES[fid]={'name':name,'book':book,'path':path,'hash':hashlib.sha256(data).hexdigest(),'converted':converted}
    sheets=[]
    for s in book:
        y,m=period(s)
        sheets.append({'name':s.title,'year':y,'month':m,'rows':s.max_row,'hidden':s.sheet_state!='visible'})
    return {'id':fid,'name':name,'sheets':sheets,'converted':converted}

def number(s,r,c):
    v=s.cell(r,c).value
    if v is None or v=='': return 0
    if isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v):
        raise ValueError(f'{s.title}!{get_column_letter(c)}{r} 不是有效數值（不支援此處的公式），請改成數值。')
    return v

def merge(req):
    t=FILES[req['template']]; wb=t['book']; out=wb[req['sheet']]; hr,cols=header(out)
    months={}
    for row in out.iter_rows(min_row=1,max_row=hr):
        for c in row:
            m=re.fullmatch(r'\s*(\d{1,2})月\s*',norm(c.value))
            if m and 1<=int(m[1])<=12:
                if int(m[1]) in months: raise ValueError('樣板含重複月份欄，請調整表頭。')
                months[int(m[1])]=c.column
    if not months: raise ValueError('樣板找不到「1月」至「12月」月份表頭。')
    ty=re.search(r'(\d{3,4})年',norm(out.cell(1,1).value)); ty=int(ty[1]) if ty else None
    targets={}
    for r in range(hr+1,out.max_row+1):
        key=(norm(out.cell(r,cols['貨物群組代碼']).value),norm(out.cell(r,cols['貨物代碼']).value))
        if not all(key): continue
        if key in targets: raise ValueError('樣板貨物代碼重複：'+'/'.join(key))
        targets[key]=r
    if not targets: raise ValueError('樣板沒有貨物資料列。')
    totals={}; seen=set(); seenhash=set(); usedmonths=set(); warnings=[]; audit=[]
    for choice in req['inputs']:
        f=FILES[choice['id']]; s=f['book'][choice['sheet']]; year=int(choice['year']); month=int(choice['month'])
        if month not in months: raise ValueError(f'{month}月不在樣板中。')
        if ty and year!=ty: raise ValueError(f'{s.title} 年度 {year} 與樣板年度 {ty} 不同。')
        identity=(f['hash'],s.title)
        if identity in seenhash: raise ValueError('同一工作表重複上傳或勾選。')
        seenhash.add(identity)
        if month in seen and not req.get('allowDuplicate'): raise ValueError(f'{month}月選取超過一張工作表，請取消重複月份，或勾選允許同月分批加總。')
        seen.add(month); usedmonths.add(month)
        h,cs=header(s)
        # TIPC 的「裝」「卸」是合併表頭，下一列的噸數才是需要加總的欄。
        tons=[c.column for c in s[h+1] if norm(c.value)=='噸數']
        if len(tons)!=2: raise ValueError(s.title+' 必須有裝、卸兩個「噸數」欄；不會猜測其他數字欄。')
        count=0; total=0
        for r in range(h+2,s.max_row+1):
            key=(norm(s.cell(r,cs['貨物群組代碼']).value),norm(s.cell(r,cs['貨物代碼']).value))
            if not all(key): continue
            val=number(s,r,tons[0])+number(s,r,tons[1]); count+=1; total+=val
            if key not in targets:
                warnings.append(f'{f["name"]}／{s.title} 第{r}列：樣板無貨物 {key[0]}/{key[1]}，未寫入 {val:,.2f} 噸。'); continue
            k=(key,month); totals[k]=totals.get(k,0)+val
        audit.append({'file':f['name'],'sheet':s.title,'month':month,'rows':count,'total':total})
    if not audit: raise ValueError('請至少選取一張輸入工作表。')
    changes={}
    for key,r in targets.items():
        for m,c in months.items():
            cell=out.cell(r,c)
            if cell.data_type=='f':
                raise ValueError(f'樣板 {cell.coordinate} 是公式，為避免覆蓋公式，請將月份明細區改為空白或數值。')
            if isinstance(cell,openpyxl.cell.cell.MergedCell): raise ValueError(f'月份資料區 {cell.coordinate} 位於合併儲存格內，無法寫入。')
            # 未匯入的月份清空，已匯入月份但無該貨物填 0。
            changes[cell.coordinate]=totals.get((key,m),0) if m in usedmonths else None
    if t['converted']: warnings.insert(0,'輸出樣板為 xls，已由 LibreOffice 轉為 xlsx。請核對轉換後的字體、合併欄位及列印設定；無法保證與原 xls 完全一致。')
    raw=patch_xlsx(t['path'].read_bytes(),wb.sheetnames.index(out.title)+1,changes)
    previewbook=openpyxl.load_workbook(io.BytesIO(raw),data_only=False)
    rid=secrets.token_hex(16); RESULTS[rid]=raw
    if len(RESULTS)>8: RESULTS.pop(next(iter(RESULTS)))
    return {'result':rid,'html':preview(previewbook[out.title]),'warnings':warnings,'audit':audit,'cells':len(changes),'months':sorted(usedmonths),'missing':sorted(set(months)-usedmonths)}

def patch_xlsx(raw,index,changes):
    # ZIP 中除目標工作表之外，所有元件逐位元保留（含樣式、列印、圖片等）。
    from xml.etree import ElementTree as ET
    zin=zipfile.ZipFile(io.BytesIO(raw)); ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    workbook=ET.fromstring(zin.read('xl/workbook.xml')); relid=list(workbook.find('m:sheets',ns))[index-1].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    rels=ET.fromstring(zin.read('xl/_rels/workbook.xml.rels')); target=next(x.get('Target') for x in rels if x.get('Id')==relid)
    import posixpath
    target=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
    xml=zin.read(target).decode('utf-8')
    for addr,value in changes.items():
        pattern=r'<c\b(?=[^>]*\br="'+addr+r'")([^>]*?)(?:/>|>(.*?)</c>)'
        match=re.search(pattern,xml,re.S)
        if match:
            attrs=re.sub(r'\s+t="[^"]*"','',match[1]); body=match[2] or ''
            body=re.sub(r'<(?:v|f|is)\b[^>]*(?:/>|>.*?</(?:v|f|is)>)','',body,flags=re.S)
            replacement='<c'+attrs+'>'+body+('' if value is None else '<v>'+str(value)+'</v>')+'</c>'
            xml=xml[:match.start()]+replacement+xml[match.end():]
        else:
            r=int(re.search(r'\d+',addr)[0]); rowpat=r'(<row\b[^>]*\br="'+str(r)+r'"[^>]*>)(.*?)(</row>)'
            rm=re.search(rowpat,xml,re.S)
            if not rm: raise ValueError('樣板缺少資料列 '+str(r))
            new='<c r="'+addr+'">'+('' if value is None else '<v>'+str(value)+'</v>')+'</c>'
            # 按欄位順序插入，以維持 Excel XML 的儲存格排序。
            from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
            wanted=column_index_from_string(coordinate_from_string(addr)[0]); body=rm[2]; at=len(body)
            for cm in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"',body):
                if column_index_from_string(cm[1])>wanted: at=cm.start(); break
            xml=xml[:rm.start(2)]+body[:at]+new+body[at:]+xml[rm.end(2):]
    dest=io.BytesIO()
    with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist(): zout.writestr(item,xml.encode('utf-8') if item.filename==target else zin.read(item.filename))
    return dest.getvalue()

def preview(s):
    def color(c,default): return '#'+c.rgb[-6:] if c and c.type=='rgb' and isinstance(c.rgb,str) else default
    merges={}; skip=set()
    for rg in s.merged_cells.ranges:
        merges[(rg.min_row,rg.min_col)]=(rg.max_row-rg.min_row+1,rg.max_col-rg.min_col+1)
        for row in rg.cells:
            if row!=(rg.min_row,rg.min_col): skip.add(row)
    widths=[]
    for col in range(1,s.max_column+1):
        letter=get_column_letter(col); dim=next((d for d in s.column_dimensions.values() if d.min and d.min<=col<=(d.max or d.min)),None)
        widths.append(max(35,((dim.width if dim else 13) or 13)*7))
    html=[f'<table class="excel" style="width:{sum(widths):.0f}px"><colgroup>']
    for col in range(1,s.max_column+1):
        dim=s.column_dimensions[get_column_letter(col)]; html.append(f'<col style="width:{widths[col-1]:.0f}px;'+('visibility:collapse' if dim.hidden else '')+'">')
    html.append('</colgroup><tbody>')
    for row in s:
        rd=s.row_dimensions[row[0].row]
        if rd.hidden: continue
        html.append(f'<tr style="height:{rd.height or 18}pt">')
        for c in row:
            if (c.row,c.column) in skip: continue
            span=merges.get((c.row,c.column),(1,1)); f=c.font; a=c.alignment
            css=f'font-family:{f.name or "Arial"};font-size:{f.sz or 11}pt;color:{color(f.color,"#111827")};background:{color(c.fill.fgColor,"#fff") if c.fill.patternType else "#fff"};text-align:{a.horizontal or ("right" if c.data_type=="n" else "left")};vertical-align:{a.vertical or "middle"};font-weight:{"bold" if f.b else "normal"};font-style:{"italic" if f.i else "normal"};white-space:{"pre-wrap" if a.wrap_text else "nowrap"};'
            for side in ('top','right','bottom','left'):
                b=getattr(c.border,side)
                if b and b.style: css+=f'border-{side}:{2 if b.style in ("medium","thick","double") else 1}px {"double" if b.style=="double" else "solid"} {color(b.color,"#444")};'
            v=c.value
            if isinstance(v,(int,float)):
                fmt=c.number_format
                dec=len(re.search(r'\.([0#]+)',fmt)[1]) if re.search(r'\.([0#]+)',fmt) else 0
                if fmt in ('General','@'): v=norm(v)
                elif '%' in fmt: v=f'{v*100:.{dec}f}%'
                else: v=format(v,(',' if ',' in fmt else '')+'.'+str(dec)+'f')
            html.append(f'<td rowspan="{span[0]}" colspan="{span[1]}" style="{escape(css)}">{escape(str(v)) if v is not None else ""}</td>')
        html.append('</tr>')
    return ''.join(html)+ '</tbody></table>'

class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args): pass
    def send(self,status,data,ctype='application/json; charset=utf-8'):
        if not isinstance(data,bytes): data=json.dumps(data,ensure_ascii=False).encode()
        self.send_response(status); self.send_header('Content-Type',ctype); self.send_header('Content-Length',str(len(data))); self.send_header('Cache-Control','no-store'); self.send_header('X-Content-Type-Options','nosniff'); self.end_headers(); self.wfile.write(data)
    def do_GET(self):
        if self.path.split('?')[0] in ('/','/2A01.html'):
            return self.send(200,(ROOT/'2A01.html').read_bytes().replace(b'__TOKEN__',TOKEN.encode()),'text/html; charset=utf-8')
        if self.path.startswith('/download/'):
            parts=self.path.split('/')
            if len(parts)==4 and parts[2]==TOKEN and parts[3] in RESULTS: return self.send(200,RESULTS[parts[3]],'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.send(404,{'error':'找不到檔案'})
    def do_POST(self):
        if self.headers.get('X-TIPC-Token')!=TOKEN: return self.send(403,{'error':'請從本機啟動頁面操作。'})
        try:
            size=int(self.headers.get('Content-Length','0'))
            if size>60*1024*1024: raise ValueError('上傳檔案過大。')
            req=json.loads(self.rfile.read(size))
            if self.path=='/upload': result=read_file(req['name'],base64.b64decode(req['data'],validate=True))
            elif self.path=='/merge': result=merge(req)
            else: return self.send(404,{'error':'找不到功能'})
            self.send(200,result)
        except Exception as e: self.send(400,{'error':str(e)})

if __name__=='__main__':
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
    url=f'http://127.0.0.1:{server.server_port}/2A01.html'
    print(url,flush=True)
    if '--no-browser' not in os.sys.argv: webbrowser.open(url)
    try: server.serve_forever()
    except KeyboardInterrupt: pass
    finally: server.server_close(); TEMP.cleanup()
